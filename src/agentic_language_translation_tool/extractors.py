"""Document extraction plugins."""

from __future__ import annotations

import re
from dataclasses import dataclass
from pathlib import Path
from typing import Any, Protocol, cast

from agentic_language_translation_tool.io import stable_id
from agentic_language_translation_tool.models import Segment

PLACEHOLDER_PATTERN = re.compile(
    r"(`[^`]+`|\{[^{}]+\}|%\([^)]+\)s|%[sd]|https?://\S+|[A-Z][A-Z0-9_]{2,})"
)
MARKDOWN_LINK_PATTERN = re.compile(r"\[([^\]]+)\]\(([^)]+)\)")


@dataclass(frozen=True)
class ExtractionResult:
    """Segments and rebuild hints extracted from a document."""

    document_format: str
    segments: list[Segment]
    structure: dict[str, object]


class Extractor(Protocol):
    """Protocol implemented by document extraction plugins."""

    format_name: str

    def supports(self, path: Path) -> bool:
        """Return whether this extractor supports a path."""

    def inspect(self, path: Path) -> dict[str, object]:
        """Return lightweight metadata for a path."""

    def extract(self, path: Path) -> ExtractionResult:
        """Extract a path into segments and structure metadata."""


class UnsupportedFormatError(ValueError):
    """Raised when no extractor exists for a file type."""


class TxtExtractor:
    """Plain-text paragraph extractor."""

    format_name = "txt"

    def supports(self, path: Path) -> bool:
        """Return whether this extractor supports TXT-like paths."""
        return path.suffix.lower() in {"", ".txt"}

    def inspect(self, path: Path) -> dict[str, object]:
        """Inspect TXT support."""
        return _base_inspection(path, self.format_name, same_format_rebuild_supported=True)

    def extract(self, path: Path) -> ExtractionResult:
        """Extract TXT paragraphs while preserving deterministic ordering."""
        text = path.read_text(encoding="utf-8")
        paragraphs = [part.strip() for part in re.split(r"\n\s*\n", text) if part.strip()]
        segments = [
            _make_segment(
                source_text=paragraph,
                document_format="txt",
                source_path=path,
                index=index,
                path=["paragraph", str(index + 1)],
                style_tags=["paragraph"],
            )
            for index, paragraph in enumerate(paragraphs)
        ]
        return ExtractionResult(
            document_format="txt",
            segments=segments,
            structure={
                "format": "txt",
                "source_file": path.name,
                "block_count": len(segments),
                "rebuild_strategy": "join_translated_paragraphs_with_blank_lines",
            },
        )


class MarkdownExtractor:
    """Markdown block-level extractor."""

    format_name = "markdown"

    def supports(self, path: Path) -> bool:
        """Return whether this extractor supports Markdown paths."""
        return path.suffix.lower() in {".md", ".markdown"}

    def inspect(self, path: Path) -> dict[str, object]:
        """Inspect Markdown support."""
        return _base_inspection(path, self.format_name, same_format_rebuild_supported=True)

    def extract(self, path: Path) -> ExtractionResult:
        """Extract Markdown block-level segments with simple formatting hints."""
        text = path.read_text(encoding="utf-8")
        blocks = _split_markdown_blocks(text)
        segments: list[Segment] = []
        structure_blocks: list[dict[str, object]] = []
        for index, block in enumerate(blocks):
            block_type = _markdown_block_type(block)
            translatable = block_type != "code_fence"
            structure_block: dict[str, object] = {
                "index": index,
                "type": block_type,
                "translatable": translatable,
                "text": block,
            }
            if not translatable:
                structure_blocks.append(structure_block)
                continue
            segment = _make_segment(
                source_text=block,
                document_format="markdown",
                source_path=path,
                index=index,
                path=[block_type, str(index + 1)],
                style_tags=[block_type],
            )
            structure_block["segment_id"] = segment.segment_id
            structure_blocks.append(structure_block)
            segments.append(segment)
        return ExtractionResult(
            document_format="markdown",
            segments=segments,
            structure={
                "format": "markdown",
                "source_file": path.name,
                "blocks": structure_blocks,
                "rebuild_strategy": "replace_translatable_blocks_by_segment_id",
            },
        )


class DocxExtractor:
    """DOCX paragraph and table-cell extractor."""

    format_name = "docx"

    def supports(self, path: Path) -> bool:
        """Return whether this extractor supports DOCX paths."""
        return path.suffix.lower() == ".docx"

    def inspect(self, path: Path) -> dict[str, object]:
        """Inspect DOCX support."""
        metadata = _base_inspection(path, self.format_name, same_format_rebuild_supported=True)
        metadata["limitations"] = [
            "Run-level formatting is preserved best-effort during rebuild.",
            "Headers, footers, comments, and footnotes are not translated in this slice.",
        ]
        return metadata

    def extract(self, path: Path) -> ExtractionResult:
        """Extract top-level paragraphs and table cells from a DOCX file."""
        from docx import Document

        document = Document(str(path))
        segments: list[Segment] = []
        blocks: list[dict[str, object]] = []
        paragraph_index = 0
        table_index = 0
        segment_index = 0

        for child in document.element.body.iterchildren():
            tag = str(child.tag)
            if tag.endswith("}p"):
                paragraph_index += 1
                paragraph = _paragraph_from_element(child, document)
                text = paragraph.text.strip()
                if not text:
                    continue
                segment_index += 1
                style_name = _safe_style_name(paragraph)
                path_parts = ["paragraph", str(paragraph_index)]
                segment = _make_segment(
                    source_text=text,
                    document_format="docx",
                    source_path=path,
                    index=segment_index,
                    path=path_parts,
                    style_tags=_docx_style_tags(style_name),
                )
                blocks.append(
                    {
                        "type": "paragraph",
                        "path": path_parts,
                        "segment_id": segment.segment_id,
                        "style": style_name,
                        "text": text,
                    }
                )
                segments.append(segment)
            elif tag.endswith("}tbl"):
                table_index += 1
                table = _table_from_element(child, document)
                for row_index, row in enumerate(table.rows, start=1):
                    for cell_index, cell in enumerate(row.cells, start=1):
                        text = "\n".join(
                            paragraph.text.strip()
                            for paragraph in cell.paragraphs
                            if paragraph.text.strip()
                        ).strip()
                        if not text:
                            continue
                        segment_index += 1
                        path_parts = [
                            "table",
                            str(table_index),
                            "row",
                            str(row_index),
                            "cell",
                            str(cell_index),
                        ]
                        segment = _make_segment(
                            source_text=text,
                            document_format="docx",
                            source_path=path,
                            index=segment_index,
                            path=path_parts,
                            style_tags=["table_cell"],
                        )
                        blocks.append(
                            {
                                "type": "table_cell",
                                "path": path_parts,
                                "segment_id": segment.segment_id,
                                "text": text,
                            }
                        )
                        segments.append(segment)

        return ExtractionResult(
            document_format="docx",
            segments=segments,
            structure={
                "format": "docx",
                "source_file": path.name,
                "blocks": blocks,
                "rebuild_strategy": "copy_source_and_replace_paragraph_or_cell_text",
                "limitations": [
                    "Run-level formatting is preserved best-effort.",
                    "Only body paragraphs and table cells are translated.",
                ],
            },
        )


class PdfExtractor:
    """Text-layer PDF block extractor."""

    format_name = "pdf"

    def supports(self, path: Path) -> bool:
        """Return whether this extractor supports PDF paths."""
        return path.suffix.lower() == ".pdf"

    def inspect(self, path: Path) -> dict[str, object]:
        """Inspect PDF support and limitations."""
        metadata = _base_inspection(path, self.format_name, same_format_rebuild_supported=False)
        metadata["limitations"] = [
            "Text-layer extraction only; OCR is not implemented.",
            "Same-format visual PDF rebuild is not supported in this slice.",
        ]
        try:
            import fitz

            with fitz.open(path) as document:
                metadata["page_count"] = document.page_count
                metadata["extractable_text_pages"] = sum(
                    1 for page in document if page.get_text("text").strip()
                )
        except Exception as error:
            metadata["inspection_warning"] = str(error)
        return metadata

    def extract(self, path: Path) -> ExtractionResult:
        """Extract text blocks from a text-layer PDF."""
        import fitz

        segments: list[Segment] = []
        pages: list[dict[str, object]] = []
        warnings: list[str] = []
        segment_index = 0
        with fitz.open(path) as document:
            for page_index, page in enumerate(document, start=1):
                page_blocks: list[dict[str, object]] = []
                text_dict = page.get_text("dict")
                raw_blocks = text_dict.get("blocks", [])
                if not isinstance(raw_blocks, list):
                    raw_blocks = []
                for block_index, block in enumerate(raw_blocks, start=1):
                    if not isinstance(block, dict) or block.get("type") != 0:
                        continue
                    text = _pdf_block_text(block)
                    if not text:
                        continue
                    segment_index += 1
                    path_parts = ["page", str(page_index), "block", str(block_index)]
                    segment = _make_segment(
                        source_text=text,
                        document_format="pdf",
                        source_path=path,
                        index=segment_index,
                        path=path_parts,
                        style_tags=["pdf_text_block"],
                    )
                    bbox = block.get("bbox", [])
                    block_info: dict[str, object] = {
                        "type": "text_block",
                        "path": path_parts,
                        "segment_id": segment.segment_id,
                        "text": text,
                        "bbox": cast(object, list(bbox) if isinstance(bbox, (list, tuple)) else []),
                    }
                    page_blocks.append(block_info)
                    segments.append(segment)
                if not page_blocks:
                    warning = f"page {page_index} has no extractable text blocks"
                    if page.get_images():
                        warning += " and may be scanned/image-only"
                    warnings.append(warning)
                pages.append({"page_number": page_index, "blocks": page_blocks})

        return ExtractionResult(
            document_format="pdf",
            segments=segments,
            structure={
                "format": "pdf",
                "source_file": path.name,
                "page_count": len(pages),
                "pages": pages,
                "warnings": warnings,
                "rebuild_strategy": "translated_markdown_by_page",
                "limitations": [
                    "Text-layer extraction only; OCR is not implemented.",
                    "Exact visual PDF recreation is not supported.",
                ],
            },
        )


EXTRACTORS: tuple[Extractor, ...] = (
    TxtExtractor(),
    MarkdownExtractor(),
    DocxExtractor(),
    PdfExtractor(),
)


def inspect_input(path: Path) -> dict[str, object]:
    """Return lightweight metadata for an input document."""
    if not path.exists():
        raise FileNotFoundError(path)
    extractor = _find_extractor(path)
    if extractor is None:
        suffix = path.suffix.lower().lstrip(".") or "unknown"
        return {
            "path": str(path),
            "format": suffix,
            "supported": False,
            "extraction_supported": False,
            "same_format_rebuild_supported": False,
            "size_bytes": path.stat().st_size,
            "supported_formats": [extractor.format_name for extractor in EXTRACTORS],
        }
    return extractor.inspect(path)


def extract_document(path: Path) -> ExtractionResult:
    """Extract supported documents into segments."""
    extractor = _find_extractor(path)
    if extractor is None:
        raise UnsupportedFormatError(f"unsupported file format: {path.suffix or '<none>'}")
    return extractor.extract(path)


def _find_extractor(path: Path) -> Extractor | None:
    for extractor in EXTRACTORS:
        if extractor.supports(path):
            return extractor
    return None


def _base_inspection(
    path: Path,
    document_format: str,
    *,
    same_format_rebuild_supported: bool,
) -> dict[str, object]:
    return {
        "path": str(path),
        "format": document_format,
        "supported": True,
        "extraction_supported": True,
        "same_format_rebuild_supported": same_format_rebuild_supported,
        "size_bytes": path.stat().st_size,
        "supported_formats": [extractor.format_name for extractor in EXTRACTORS],
    }


def extract_txt(text: str, source_path: Path) -> ExtractionResult:
    """Compatibility wrapper for TXT extraction tests and callers."""
    return TxtExtractor().extract(_write_virtual_text(source_path, text))


def extract_markdown(text: str, source_path: Path) -> ExtractionResult:
    """Compatibility wrapper for Markdown extraction tests and callers."""
    return MarkdownExtractor().extract(_write_virtual_text(source_path, text))


def _write_virtual_text(source_path: Path, text: str) -> Path:
    if source_path.exists():
        return source_path
    source_path.parent.mkdir(parents=True, exist_ok=True)
    source_path.write_text(text, encoding="utf-8")
    return source_path


def _split_markdown_blocks(text: str) -> list[str]:
    blocks: list[str] = []
    current: list[str] = []
    in_fence = False
    for line in text.splitlines():
        if line.strip().startswith("```"):
            current.append(line)
            if in_fence:
                blocks.append("\n".join(current).strip())
                current = []
            in_fence = not in_fence
            continue
        if in_fence:
            current.append(line)
            continue
        if not line.strip():
            if current:
                blocks.append("\n".join(current).strip())
                current = []
            continue
        current.append(line)
    if current:
        blocks.append("\n".join(current).strip())
    return blocks


def _markdown_block_type(block: str) -> str:
    stripped = block.lstrip()
    if stripped.startswith("```"):
        return "code_fence"
    if stripped.startswith("#"):
        return "heading"
    if stripped.startswith(("- ", "* ", "+ ")) or re.match(r"^\d+\.\s", stripped):
        return "list"
    if "|" in stripped and "\n" in stripped:
        return "table"
    if MARKDOWN_LINK_PATTERN.search(stripped):
        return "linked_paragraph"
    return "paragraph"


def _make_segment(
    *,
    source_text: str,
    document_format: str,
    source_path: Path,
    index: int,
    path: list[str],
    style_tags: list[str],
) -> Segment:
    placeholders = sorted(set(PLACEHOLDER_PATTERN.findall(source_text)))
    protected_terms = sorted(term for term in placeholders if _looks_protected(term))
    checksum = stable_id(source_text, length=32)
    segment_id = f"seg_{stable_id(str(source_path), str(index), source_text)}"
    return Segment(
        segment_id=segment_id,
        source_text=source_text,
        format=document_format,
        context=f"{source_path.name} block {index}",
        path=path,
        style_tags=style_tags,
        placeholders=placeholders,
        protected_terms=protected_terms,
        checksum=checksum,
    )


def _paragraph_from_element(element: Any, document: Any) -> Any:
    from docx.text.paragraph import Paragraph

    return Paragraph(element, document)


def _table_from_element(element: Any, document: Any) -> Any:
    from docx.table import Table

    return Table(element, document)


def _safe_style_name(paragraph: Any) -> str:
    style = getattr(paragraph, "style", None)
    name = getattr(style, "name", None)
    return str(name) if name else "Normal"


def _docx_style_tags(style_name: str) -> list[str]:
    tags = [style_name]
    normalized = style_name.lower()
    if normalized.startswith("heading"):
        tags.append("heading")
    if "list" in normalized:
        tags.append("list")
    return tags


def _pdf_block_text(block: dict[str, object]) -> str:
    lines = block.get("lines", [])
    if not isinstance(lines, list):
        return ""
    rendered_lines: list[str] = []
    for line in lines:
        if not isinstance(line, dict):
            continue
        spans = line.get("spans", [])
        if not isinstance(spans, list):
            continue
        line_text = "".join(
            str(span.get("text", ""))
            for span in spans
            if isinstance(span, dict)
        ).strip()
        if line_text:
            rendered_lines.append(line_text)
    return "\n".join(rendered_lines).strip()


def _looks_protected(value: str) -> bool:
    return value.startswith(("http://", "https://", "`")) or value.isupper()
